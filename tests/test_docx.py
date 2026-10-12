"""Regression checks for Word sizing, direction, and protected content."""
import importlib.util
from pathlib import Path
import tempfile
import unittest
import unicodedata

ROOT = Path(__file__).resolve().parents[1]
DOCX_AVAILABLE = importlib.util.find_spec("docx") is not None


@unittest.skipUnless(DOCX_AVAILABLE, "Word fixture checks require python-docx")
class WordFixtureTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from docx import Document
        from docx.oxml.ns import qn
        cls.qn = staticmethod(qn)
        cls.temp = tempfile.TemporaryDirectory()
        cls.addClassCleanup(cls.temp.cleanup)
        spec = importlib.util.spec_from_file_location("word_builder", ROOT / "examples" / "build_docx.py")
        builder = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(builder)
        output = Path(cls.temp.name) / "sample.docx"
        builder.build(output)
        cls.doc = Document(output)

    def test_title_keeps_style_size(self):
        title = self.doc.paragraphs[0]
        self.assertEqual(title.style.name, "Title")
        self.assertIsNone(title.runs[0].font.size)
        self.assertIsNone(title.runs[0]._r.rPr.find(self.qn("w:szCs")))
        style = title.style
        self.assertGreater(style.font.size.pt, 12)
        self.assertEqual(int(style.element.rPr.find(self.qn("w:szCs")).get(self.qn("w:val"))),
                         int(style.font.size.pt * 2))
        self.assertIsNone(style.element.pPr.find(self.qn("w:pBdr")))

    def test_paragraph_alignment_follows_direction(self):
        for paragraph in self.doc.paragraphs:
            self.assertEqual(paragraph._p.pPr.find(self.qn("w:jc")).get(self.qn("w:val")), "start")
        code = self.doc.paragraphs[-1]
        self.assertEqual(code._p.pPr.find(self.qn("w:bidi")).get(self.qn("w:val")), "0")

    def test_date_uses_native_ltr_cell(self):
        row = next(row for row in self.doc.tables[0].rows if row.cells[0].text == "تاریخ آزمون")
        self.assertEqual(row.cells[1].text, "2026-10-12")
        self.assertEqual(row.cells[1].paragraphs[0]._p.pPr.find(self.qn("w:bidi")).get(self.qn("w:val")), "0")
        self.assertEqual(row.cells[1].paragraphs[0].runs[0]._r.rPr.find(self.qn("w:rtl")).get(self.qn("w:val")), "0")

    def test_technical_content_and_size(self):
        texts = self.doc.element.body.xpath(".//w:t/text()")
        for token in ["v2.1.0", "8080", "2026-10-12", "3.14",
                      r"C:\Projects\demo\config.json", "https://example.com/docs?q=rtl&v=2"]:
            self.assertIn(token, texts)
        self.assertFalse(any(unicodedata.category(ch) == "Cf" and ch != "\u200c"
                             for text in texts for ch in text))
        self.assertEqual(self.doc.paragraphs[1].runs[0].font.size.pt, 12)
        code = self.doc.paragraphs[-1].runs[0]
        self.assertEqual(code.text, 'payload = {"version": "v2.1.0", "amount": 3.14}')
        self.assertEqual(code.font.size.pt, 10)
        self.assertEqual(code._r.rPr.find(self.qn("w:szCs")).get(self.qn("w:val")), "20")
        targets = [rel.target_ref for rel in self.doc.part.rels.values() if rel.is_external]
        self.assertIn("https://example.com/docs?q=rtl&v=2", targets)


if __name__ == "__main__":
    unittest.main()
