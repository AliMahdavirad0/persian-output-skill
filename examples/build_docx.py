"""Create an RTL fixture. Requires python-docx; render separately before acceptance."""
import argparse
from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE as RT


def property_node(props, tag, value="1"):
    node = props.find(qn(tag))
    if node is None:
        node = OxmlElement(tag)
        props.append(node)
    node.set(qn("w:val"), value)
    return node


def direction(paragraph, rtl=True):
    property_node(paragraph._p.get_or_add_pPr(), "w:bidi", "1" if rtl else "0")
    # Use logical alignment: start is right for RTL and left for LTR.
    property_node(paragraph._p.get_or_add_pPr(), "w:jc", "start")


def run_text(paragraph, text, rtl=True, size=None):
    run = paragraph.add_run(text)
    run.font.name = "Tahoma" if rtl else "Consolas"
    if size is not None:
        run.font.size = Pt(size)
    props = run._r.get_or_add_rPr()
    property_node(props, "w:rtl", "1" if rtl else "0")
    fonts = props.find(qn("w:rFonts"))
    fonts.set(qn("w:cs"), "Tahoma" if rtl else "Consolas")
    if size is not None:
        property_node(props, "w:szCs", str(int(size * 2)))
    lang = property_node(props, "w:lang", "fa-IR" if rtl else "en-US")
    lang.set(qn("w:bidi"), "fa-IR" if rtl else "en-US")
    return run


def paragraph(doc, segments, style=None, rtl=True):
    p = doc.add_paragraph(style=style)
    direction(p, rtl)
    for text, is_rtl in segments:
        size = None if style and style != "List Bullet" else (12 if is_rtl else 10)
        run_text(p, text, is_rtl, size=size)
    return p


def build(output):
    doc = Document()
    # Match complex-script title size to the existing Title style, not body text.
    title_style = doc.styles["Title"]
    property_node(title_style.element.get_or_add_rPr(), "w:szCs",
                  str(int(title_style.font.size.pt * 2)))
    title_style.font.color.rgb = RGBColor(0, 0, 0)
    borders = title_style.element.get_or_add_pPr().find(qn("w:pBdr"))
    if borders is not None:
        borders.getparent().remove(borders)
    doc.core_properties.author = "Persian Output Skill"
    doc.core_properties.title = "نمونهٔ خروجی فارسی"
    for section in doc.sections:
        section.top_margin = section.bottom_margin = Inches(0.8)
        section.left_margin = section.right_margin = Inches(0.8)
    paragraph(doc, [("نمونهٔ خروجی فارسی", True)], "Title")
    paragraph(doc, [("این نمونه برای بررسی جهت بندها، جدول و متن ترکیبی است. تأیید نمایش به رندر و مشاهدهٔ صفحه نیاز دارد.", True)])
    paragraph(doc, [("این فایل‌ها به‌روز می‌شوند؛ آیا نیم‌فاصله‌ها درست هستند؟", True)])
    paragraph(doc, [("نسخهٔ ", True), ("v2.1.0", False), (" (نسخهٔ آزمایشی) روی پورت ", True), ("8080", False), (" اجرا می‌شود.", True)])
    paragraph(doc, [("مسیر فایل: ", True), (r"C:\Projects\demo\config.json", False)])
    for text in ["معنای متن حفظ شود.", "کد و لینک‌ها دقیق باقی بمانند."]:
        paragraph(doc, [(text, True)], "List Bullet")
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    property_node(table._tbl.tblPr, "w:bidiVisual")
    for cell, value in zip(table.rows[0].cells, ["مورد", "مقدار"]):
        direction(cell.paragraphs[0])
        run_text(cell.paragraphs[0], value, size=12)
    # Numeric dates have only weak bidi characters. Give technical values a
    # native LTR cell paragraph so renderers keep their order without marks.
    rows = [
        ("نسخه", "v2.1.0", False),
        ("تاریخ آزمون", "2026-10-12", False),
        ("مقدار فنی", "3.14", False),
        ("تعداد نمونه‌ها", "۱۲", True),
        ("وضعیت", "آمادهٔ بررسی", True),
    ]
    for label, value, is_rtl in rows:
        cells = table.add_row().cells
        direction(cells[0].paragraphs[0])
        direction(cells[1].paragraphs[0], rtl=is_rtl)
        run_text(cells[0].paragraphs[0], label, size=12)
        run_text(cells[1].paragraphs[0], value, is_rtl, size=12 if is_rtl else 10)
    p = paragraph(doc, [("پیوند: ", True)])
    p.paragraph_format.space_before = Pt(8)
    target = "https://example.com/docs?q=rtl&v=2"
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), p.part.relate_to(target, RT.HYPERLINK, is_external=True))
    r = run_text(p, target, False, size=10)
    p._p.remove(r._r)
    link.append(r._r)
    p._p.append(link)
    paragraph(doc, [('payload = {"version": "v2.1.0", "amount": 3.14}', False)], rtl=False)
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    print(f"Created {output}; visual review still required.")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", required=True, type=Path)
    build(parser.parse_args().output)
